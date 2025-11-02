from docx import Document
from docx.shared import Pt

# Metadata (from user)
author_name = "Rtamanyu N.J."
roll = "24f3001383"
email = "24f3001383@ds.study.iitm.ac.in"
about = "I will add later while editing the doc"
video_link = "I will provide while editing the doc"

# Create document
doc = Document()

# Set default font size
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Title
doc.add_heading('Project Report - Hospital Management System', level=1)
# Project Details (matches example format)
doc.add_heading('Project Details', level=2)
doc.add_paragraph('Project Title: Hospital Management Web App', style='Intense Quote')
doc.add_paragraph('Problem Statement:', style='Heading 3')
doc.add_paragraph('To design and build a web-based application that allows administrators, doctors and patients to manage hospital workflows — user accounts, appointments scheduling, department and doctor management, and simple treatment records. The application should provide calendar-based availability, prevent double-bookings, and expose a compact public API for availability checks.')
doc.add_paragraph('Approach:', style='Heading 3')
doc.add_paragraph('The app was built using Flask with a modular structure. Server-side rendering (Jinja2) provides baseline UX; richer interactions are implemented with externalized JavaScript in `static/js/`. A REST API (Flask-RESTful) complements the UI for AJAX interactions and the public availability endpoint. SQLAlchemy ORM manages persistence to a local SQLite database.')

# Author block
doc.add_heading('Author', level=2)
doc.add_paragraph(f'Name: {author_name}')
doc.add_paragraph(f'Roll No.: {roll}')
doc.add_paragraph(f'Email: {email}')
doc.add_paragraph(f'About: {about}')

# AI/LLM usage
doc.add_heading('AI / LLM usage', level=2)
doc.add_paragraph('Used an LLM (ChatGPT) to assist with code refactoring, externalizing JavaScript, designing and implementing a REST API for appointments (including a public availability endpoint), and adding server- and client-side validation. The LLM was used for suggestions, iterative code generation, and creating documentation snippets. Final code changes were reviewed and integrated by the developer.')

# Description
doc.add_heading('Description', level=2)
doc.add_paragraph('A minimal Hospital Management web application built with Flask that supports three roles (admin, doctor, patient).')
doc.add_paragraph('Detailed description:')
doc.add_paragraph('This application allows patients to register and manage their profile, browse available doctors by department, view doctors\' availability on a calendar-style interface, and book 30-minute appointment slots. Doctors can view their upcoming schedule, see basic patient details for booked slots, mark appointments as completed, and add treatment notes. Administrators can create and manage departments, add/edit doctor and patient profiles, and blacklist accounts when required.')
doc.add_paragraph('The booking engine ensures that no two non-cancelled appointments exist for the same doctor/date/time combination. Avatars are stored as binary blobs in the database and served via dedicated avatar endpoints; a default placeholder image is returned when an avatar is missing. The frontend follows a progressive enhancement approach: core interactions work with server-rendered pages while richer experiences (dynamic appointment refresh, calendar UI) are powered by externalized JavaScript in `static/js/` that consumes JSON endpoints.')

# Technologies and Frameworks Used
doc.add_heading('Technologies and Frameworks Used', level=2)
# Create a 2-column table: Technology / Purpose
tech_table = doc.add_table(rows=1, cols=2)
hdr_cells = tech_table.rows[0].cells
hdr_cells[0].text = 'Technology / Library'
hdr_cells[1].text = 'Purpose'
rows = [
	('Flask', 'Core backend web framework'),
	('SQLAlchemy', 'Object Relational Mapper for SQLite database'),
	('Jinja2', 'Template engine for rendering dynamic HTML pages'),
	('Bootstrap 5', 'Frontend styling and responsive design'),
	('Flask-Login', 'User authentication and session management'),
	('Flask-RESTful', 'Lightweight REST API framework'),
	('SQLite', 'Lightweight local database for storing user data')
]
for tech, purpose in rows:
	r = tech_table.add_row().cells
	r[0].text = tech
	r[1].text = purpose

# DB Schema summary
doc.add_heading('DB Schema Design', level=2)
doc.add_paragraph('Core tables: Login, Patient, Doctor, Department, Appointment, Treatment. See models.py for exact columns. Important constraints: Appointment uniqueness per doctor/date/time enforced at application level; foreign keys link profiles to Login entries.')

# API Design
doc.add_heading('API Design', level=2)
doc.add_paragraph('Implemented REST API endpoints using Flask-RESTful. A public endpoint exposes doctor availability (booked/free) over a date range; authenticated endpoints support listing, creating, retrieving, rescheduling and cancelling appointments. An OpenAPI/YAML file should be included separately with the final submission.')

# Architecture and Features
doc.add_heading('Architecture Overview', level=2)
doc.add_paragraph('The project is organized into clear components; key files and folders:')
doc.add_paragraph('main.py – main Flask application entry point', style='List Bullet')
doc.add_paragraph('application/models.py – database models using SQLAlchemy', style='List Bullet')
doc.add_paragraph('application/controllers.py – Flask routes and controller logic', style='List Bullet')
doc.add_paragraph('application/api_appointments.py – RESTful API resources for appointments', style='List Bullet')
doc.add_paragraph('templates/ – Jinja2 HTML templates', style='List Bullet')
doc.add_paragraph('static/ – CSS, JS and client assets (calendar, AJAX handlers)', style='List Bullet')

doc.add_heading('Implemented Features', level=2)
features = [
	'User registration and login',
	'Daily appointment booking and calendar availability',
	'Doctor dashboard with appointment management',
	'Admin management for doctors, patients and departments',
	'Profile avatars with upload and fallback handling',
	'Input validation (client and server) and blacklist controls'
]
for f in features:
	doc.add_paragraph(f, style='List Bullet')

doc.add_heading('Additional Features', level=3)
doc.add_paragraph('Export user logs as CSV', style='List Bullet')
doc.add_paragraph('AI/LLM-generated weekly summary (optional extension)', style='List Bullet')
doc.add_paragraph('Public availability API for external schedulers', style='List Bullet')

doc.add_paragraph('')
doc.add_heading('Detailed Features', level=2)
doc.add_paragraph('1) Registration and Authentication:', style='List Number')
doc.add_paragraph('Users register using an email (used as username). Password strength checks are enforced on both client and server. Flask-Login manages session authentication. Admins, doctors and patients have distinct dashboards and route protections.')
doc.add_paragraph('2) Appointment Booking & Availability:', style='List Number')
doc.add_paragraph('Doctors expose 30-minute slots between 09:00 and 20:00. Patients can view a 7-day default window (configurable) and book free slots. Server-side checks prevent double-booking and validate date/time formats. A public API endpoint exposes booked/free slots for a doctor for integration with external systems.')
doc.add_paragraph('3) Doctor & Admin Workflows:', style='List Number')
doc.add_paragraph('Doctors can view their schedule, mark appointments as completed, and add treatment records linked to appointments. Admins can CRUD doctors, patients and departments. Deletions are guarded (e.g., a department cannot be deleted if doctors exist).')
doc.add_paragraph('4) Avatar and Media Handling:', style='List Number')
doc.add_paragraph('Profile pictures are uploaded as multipart form data and stored as BLOBs. Avatar endpoints detect image type and return appropriate content type; missing avatars redirect to a default asset. Cache-busting is performed on client-side when profile images are updated to ensure freshness.')
doc.add_paragraph('5) Validation & Security:', style='List Number')
doc.add_paragraph('Input validation (email, phone, password strength) is implemented on both client and server. SQL injection risk is mitigated by using SQLAlchemy ORM. Blacklist functionality allows admins to disable accounts. Passwords are currently stored in plaintext in the DB schema and should be migrated to hashed storage before any production use.')
doc.add_paragraph('6) REST API:', style='List Number')
doc.add_paragraph('A set of Flask-RESTful endpoints provide JSON access to appointment data. Authenticated endpoints require session login and enforce role-based access. A public availability endpoint returns a compact availability structure for third-party integrations.')
doc.add_paragraph('7) UX & Frontend:', style='List Number')
doc.add_paragraph('Inline JavaScript was removed from templates and externalized to `static/js/` files. The frontend uses Bootstrap 5 for layout and responsiveness. The patient dashboard contains tabs for Upcoming and Past appointments which are updated dynamically via AJAX calls to the server.')
doc.add_paragraph('Future improvements: move to hashed passwords, add API tokens or OAuth for programmatic clients, add pagination and filtering to APIs, and include automated tests and CI/CD for deployments.')

doc.add_heading('Data Flow & Internals', level=2)
doc.add_paragraph('1) Request flow: A typical booking request originates from the patient dashboard where JavaScript sends a JSON POST to the server (or a form submit). The server validates the payload (doctor_id, date, time), checks slot availability against the Appointment table and, on success, inserts a new Appointment row and returns a minimal JSON response. The page then updates the UI via AJAX callbacks.')
doc.add_paragraph('2) Avatar handling: Profile images are sent as multipart/form-data, read into memory and stored in BLOB fields on the related profile table. Avatar endpoints stream the bytes back with the proper Content-Type (image/png, image/jpeg) using simple content detection. When no avatar exists, a redirect to a default static image is returned, allowing CDN-friendly caching of the default asset.')

doc.add_heading('Testing, Validation & Error Handling', level=2)
doc.add_paragraph('Unit and integration tests are recommended but not included in this submission. The server-side validates all inputs (email, phone, date/time formats) and returns appropriate HTTP status codes (400 for bad requests, 403 for unauthorized, 404 for not found, 409 for conflicts). Client-side validation mirrors these checks to provide immediate feedback and reduce round trips. Database errors are caught and rolled back with clear JSON error messages for AJAX endpoints and flash messages for HTML workflows.')

doc.add_heading('Deployment & Environment', level=2)
doc.add_paragraph('The app is lightweight and can be served with a WSGI server (gunicorn/uWSGI) behind a reverse proxy. For production readiness:')
doc.add_paragraph('- Move secret config (SECRET_KEY, DB path) into environment variables or a secure vault.')
doc.add_paragraph('- Replace plaintext passwords with a salted hash (bcrypt/argon2) and enforce strong password rules.')
doc.add_paragraph('- Use a managed database for scaling; update storage strategy for profile images (e.g., object storage instead of BLOBs).')

doc.add_heading('Performance & Limitations', level=2)
doc.add_paragraph('This prototype is intended for demonstration and testing. Limitations:')
doc.add_paragraph('- No pagination on list endpoints; large datasets may slow responses.')
doc.add_paragraph('- Images are stored in the DB which may impact DB performance; consider external storage for production.')
doc.add_paragraph('- No rate limiting or API keys for the public endpoint; consider adding rate limiting and caching for high-volume traffic.')

doc.add_heading('Sample usage of the public availability API', level=2)
doc.add_paragraph('Example curl command to retrieve availability for doctor with id=2 for the next 3 days starting 2025-11-05:')
doc.add_paragraph('curl "http://127.0.0.1:8000/api/public/doctor/2/availability?start=2025-11-05&days=3"')

doc.add_heading('How to run (short)', level=2)
doc.add_paragraph('1) Create and activate the virtual environment in the project root')
doc.add_paragraph('2) Install dependencies: pip install -r requirements.txt (or pip install flask flask-login sqlalchemy flask-restful python-docx)')
doc.add_paragraph('3) Initialize DB (if needed) and run: python main.py (development server binds to 0.0.0.0:8000 by default)')

doc.add_heading('Known issues and next steps', level=2)
doc.add_paragraph('1) Migrate passwords to hashed storage. 2) Add automated tests and CI. 3) Add OpenAPI YAML for all endpoints and include a small test harness for the public API. 4) Add rate limiting and monitoring for the public endpoint.')

# Video
doc.add_heading('Video', level=2)
doc.add_paragraph(f'Video link: {video_link}')

# Final submission notes
doc.add_heading('Final submission', level=2)
doc.add_paragraph('The final ZIP should include: report.pdf (2-5 pages), the code folder with a README explaining how to run the project, an api.yaml file describing the API, and a small text file with the video link.')

# Footer / contact
p = doc.add_paragraph()
p.add_run('Prepared by: ').bold = True
p.add_run(author_name + ' ').italic = True
p.add_run(f'| Roll No.: {roll} | Email: {email}')

# Save
out_path = 'Project_Report_Rtamanyu_NJ.docx'
doc.save(out_path)
print('Saved:', out_path)
